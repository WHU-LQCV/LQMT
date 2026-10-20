"""
结果导出模块
支持导出Excel、CSV、PDF格式
"""

import pandas as pd
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

try:
    from .utils import rename_metric_columns
except ImportError:
    from gui_app.utils import rename_metric_columns


class ResultExporter:
    """结果导出器"""
    
    def __init__(self):
        """初始化导出器"""
        self.grade_colors = {
            'A': 'C6EFCE',  # 浅绿色
            'B': 'C6E0F2',  # 浅蓝色
            'C': 'FFE699',  # 浅黄色
            'D': 'FFC7CE',  # 浅红色
        }
    
    def _reorder_columns(self, results_df, selected_metrics=None):
        """重新排序列以满足新的展示顺序，并根据选中的指标筛选"""
        # 简化列名到完整列名的映射
        short_to_full = {
            'Ra': 'Ra (CIE 13.3-1995 Ra)',
            'Rf': 'Rf (CIE 224-2017 Rf / IES TM 30-18 Rf)',
            'Rg': 'Rg (IES TM30-18 Rg)',
        }
        
        # 完整列名到简化列名的映射
        full_to_short = {v: k for k, v in short_to_full.items()}
        
        # 定义列的顺序（使用简化名称）
        ordered_short_names = ['ID', '还原分级', '喜好分级', '分辨分级',
                              'Ra', 'Rf', '(Ra+Rf)/2', 'MCPI', 'CDM', 'MCPIr', 'CDMr',
                              'CCT', 'Duv', 'x', 'y', "u'", "v'", 'pe', 'hab', 
                              'Rg', 'CRI2012', 
                              'CQS Qa', 'CQS Qf', 'CQS Qg', 'CQS Qp',
                              'MCRI', 'FCI', 'Sneutral', 'GAI', 'CDI76', 'GAI-CRI',
                              'Rd', 'CS', 'CLa']
        
        # 构建最终的列顺序
        final_columns = []
        for short_name in ordered_short_names:
            # 检查是否被选中（如果有selected_metrics的话）
            if selected_metrics is not None and short_name != 'ID':
                if not selected_metrics.get(short_name, True):
                    continue
            
            # 获取完整列名
            full_name = short_to_full.get(short_name, short_name)
            
            if full_name in results_df.columns:
                final_columns.append(full_name)
        
        # 添加其他未在列表中的列
        for col in results_df.columns:
            if col not in final_columns:
                # 获取简化名称检查是否被选中
                short_name = full_to_short.get(col, col)
                if selected_metrics is not None:
                    if not selected_metrics.get(short_name, True):
                        continue
                final_columns.append(col)
        
        return results_df[final_columns]
    
    def export_to_excel(self, results_df, spd_array=None, sample_names=None, 
                       file_path='output.xlsx', selected_metrics=None):
        """
        导出到Excel文件（多Sheet + 格式化）
        
        参数:
            results_df: DataFrame - 计算结果
            spd_array: ndarray - SPD数组
            sample_names: list - 样本名称
            file_path: str - 输出文件路径
            selected_metrics: dict - 选中的指标字典
        """
        # 重新排序列并筛选。摘要仍用内部列名，写出的表头换成下划线名称。
        results_df = self._reorder_columns(results_df, selected_metrics)
        display_df = rename_metric_columns(results_df)
        
        with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
            # Sheet 1: 摘要信息
            self._write_summary_sheet(writer, results_df)
            
            # Sheet 2: 完整指标数据
            self._write_metrics_sheet(writer, display_df, apply_formatting=True)
            
            # Sheet 3: SPD原始数据（如果提供）
            if spd_array is not None and sample_names is not None:
                self._write_spd_sheet(writer, spd_array, sample_names)
            
            # Sheet 4: AAA级光源（如果存在，且本次计算包含分级列）
            grade_cols = ['还原分级', '喜好分级', '分辨分级']
            if all(col in results_df.columns for col in grade_cols):
                aaa_df = results_df[
                    (results_df['还原分级'] == 'A') &
                    (results_df['喜好分级'] == 'A') &
                    (results_df['分辨分级'] == 'A')
                ]
                if len(aaa_df) > 0:
                    aaa_df = rename_metric_columns(aaa_df)
                    aaa_df.to_excel(writer, sheet_name='AAA级光源', index=False)
                    self._apply_sheet_formatting(writer.sheets['AAA级光源'])
    
    def _write_summary_sheet(self, writer, results_df):
        """写入摘要信息Sheet"""
        summary_rows = [
            ('总样本数', len(results_df)),
        ]

        if 'CCT' in results_df.columns:
            summary_rows.extend([
                ('平均CCT (K)', int(results_df['CCT'].mean())),
                ('CCT范围', f"{results_df['CCT'].min()}K - {results_df['CCT'].max()}K"),
            ])
        if 'Ra (CIE 13.3-1995 Ra)' in results_df.columns:
            summary_rows.append(('平均 R_a', round(results_df['Ra (CIE 13.3-1995 Ra)'].mean(), 1)))
        if 'Rf (CIE 224-2017 Rf / IES TM 30-18 Rf)' in results_df.columns:
            summary_rows.append(('平均 R_f', round(results_df['Rf (CIE 224-2017 Rf / IES TM 30-18 Rf)'].mean(), 1)))

        grade_cols = ['还原分级', '喜好分级', '分辨分级']
        if all(col in results_df.columns for col in grade_cols):
            summary_rows.extend([
                ('AAA级样本数', len(results_df[
                    (results_df['还原分级'] == 'A') &
                    (results_df['喜好分级'] == 'A') &
                    (results_df['分辨分级'] == 'A')
                ])),
                ('B级及以上（还原）', len(results_df[results_df['还原分级'].isin(['A', 'B'])])),
                ('B级及以上（喜好）', len(results_df[results_df['喜好分级'].isin(['A', 'B'])])),
                ('B级及以上（分辨）', len(results_df[results_df['分辨分级'].isin(['A', 'B'])])),
            ])

        summary_df = pd.DataFrame(summary_rows, columns=['项目', '数值'])
        summary_df.to_excel(writer, sheet_name='摘要信息', index=False)
        
        # 应用格式
        ws = writer.sheets['摘要信息']
        self._apply_summary_formatting(ws)
    
    def _write_metrics_sheet(self, writer, results_df, apply_formatting):
        """写入完整指标数据Sheet"""
        results_df.to_excel(writer, sheet_name='完整指标数据', index=False)
        
        if apply_formatting:
            ws = writer.sheets['完整指标数据']
            self._apply_metrics_formatting(ws, results_df)
    
    def _write_spd_sheet(self, writer, spd_array, sample_names):
        """写入SPD数据Sheet"""
        # 构建DataFrame
        spd_dict = {'Wavelength': spd_array[0, :].tolist()}
        for i, name in enumerate(sample_names):
            spd_dict[name] = spd_array[i + 1, :].tolist()
        
        spd_df = pd.DataFrame(spd_dict)
        spd_df.to_excel(writer, sheet_name='SPD数据', index=False)
        
        # 应用格式
        ws = writer.sheets['SPD数据']
        self._apply_sheet_formatting(ws)
    
    def _apply_summary_formatting(self, ws):
        """应用摘要信息格式"""
        # 设置列宽
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 20
        
        # 标题行格式
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF', size=12)
        
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        # 数据行格式
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.alignment = Alignment(horizontal='left', vertical='center')
                cell.border = Border(
                    left=Side(style='thin'),
                    right=Side(style='thin'),
                    top=Side(style='thin'),
                    bottom=Side(style='thin')
                )
    
    def _apply_metrics_formatting(self, ws, results_df):
        """应用指标数据格式"""
        # 设置列宽
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 30)
            ws.column_dimensions[column].width = adjusted_width
        
        # 标题行格式
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF', size=11)
        
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        # 查找分级列的位置
        header_row = [cell.value for cell in ws[1]]
        fidelity_col = header_row.index('还原分级') + 1 if '还原分级' in header_row else None
        preference_col = header_row.index('喜好分级') + 1 if '喜好分级' in header_row else None
        discrimination_col = header_row.index('分辨分级') + 1 if '分辨分级' in header_row else None
        
        # 应用分级颜色
        for row_idx, (_, row_data) in enumerate(results_df.iterrows(), start=2):
            # 还原分级
            if fidelity_col:
                cell = ws.cell(row=row_idx, column=fidelity_col)
                grade = row_data['还原分级']
                if grade in self.grade_colors:
                    cell.fill = PatternFill(start_color=self.grade_colors[grade], 
                                          end_color=self.grade_colors[grade], 
                                          fill_type='solid')
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal='center')
            
            # 喜好分级
            if preference_col:
                cell = ws.cell(row=row_idx, column=preference_col)
                grade = row_data['喜好分级']
                if grade in self.grade_colors:
                    cell.fill = PatternFill(start_color=self.grade_colors[grade], 
                                          end_color=self.grade_colors[grade], 
                                          fill_type='solid')
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal='center')
            
            # 分辨分级
            if discrimination_col:
                cell = ws.cell(row=row_idx, column=discrimination_col)
                grade = row_data['分辨分级']
                if grade in self.grade_colors:
                    cell.fill = PatternFill(start_color=self.grade_colors[grade], 
                                          end_color=self.grade_colors[grade], 
                                          fill_type='solid')
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal='center')
        
        # 添加边框
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
            for cell in row:
                cell.border = thin_border
    
    def _apply_sheet_formatting(self, ws):
        """应用通用Sheet格式"""
        # 设置列宽
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 25)
            ws.column_dimensions[column].width = adjusted_width
        
        # 标题行格式
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF', size=11)
        
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        # 添加边框
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
            for cell in row:
                cell.border = thin_border
    
    def _add_charts_to_excel(self, file_path, results_df, spd_array, sample_names):
        """向Excel添加图表（可选功能）"""
        # 这里可以使用openpyxl的图表功能或者matplotlib生成图片插入
        # 由于复杂度较高，这里提供基础框架
        try:
            from openpyxl.chart import BarChart, Reference
            
            wb = openpyxl.load_workbook(file_path)
            
            # 在摘要信息Sheet添加分级统计图表
            if '摘要信息' in wb.sheetnames:
                ws = wb['摘要信息']
                # TODO: 添加图表逻辑
            
            wb.save(file_path)
        except Exception as e:
            print(f"添加图表失败: {e}")
    
    def export_to_csv(self, results_df, file_path='output.csv', selected_metrics=None):
        """
        导出到CSV文件
        
        参数:
            results_df: DataFrame - 计算结果
            file_path: str - 输出文件路径
            selected_metrics: dict - 选中的指标字典
        """
        # 重新排序列并筛选
        results_df = rename_metric_columns(self._reorder_columns(results_df, selected_metrics))
        # 使用utf-8-sig编码，确保Excel正确显示中文
        results_df.to_csv(file_path, index=False, encoding='utf-8-sig')
    
    def export_to_pdf(self, results_df, file_path='output.pdf'):
        """
        导出到PDF报告（使用reportlab）
        
        参数:
            results_df: DataFrame - 计算结果
            file_path: str - 输出文件路径
        """
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import letter, A4
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import inch
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
            
            # 注册中文字体（需要系统有相应字体文件）
            # pdfmetrics.registerFont(TTFont('SimSun', 'simsun.ttc'))
            
            doc = SimpleDocTemplate(file_path, pagesize=A4)
            elements = []
            
            styles = getSampleStyleSheet()
            title_style = styles['Title']
            
            # 标题
            title = Paragraph("LQCV 光品质指标计算报告", title_style)
            elements.append(title)
            elements.append(Spacer(1, 0.5*inch))
            
            # 摘要信息
            summary_items = [f"<b>总样本数:</b> {len(results_df)}"]
            if 'CCT' in results_df.columns:
                summary_items.append(f"<b>平均CCT:</b> {int(results_df['CCT'].mean())}K")
            grade_cols = ['还原分级', '喜好分级', '分辨分级']
            if all(col in results_df.columns for col in grade_cols):
                summary_items.append(
                    "<b>AAA级样本数:</b> " +
                    str(len(results_df[
                        (results_df['还原分级'] == 'A') &
                        (results_df['喜好分级'] == 'A') &
                        (results_df['分辨分级'] == 'A')
                    ]))
                )
            summary_text = "<br/>".join(summary_items)
            summary = Paragraph(summary_text, styles['Normal'])
            elements.append(summary)
            elements.append(Spacer(1, 0.3*inch))
            
            # 创建表格（显示前10行）
            table_data = [results_df.columns.tolist()]
            for _, row in results_df.head(10).iterrows():
                table_data.append(row.tolist())
            
            table = Table(table_data)
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))
            
            elements.append(table)
            
            doc.build(elements)
            
        except ImportError:
            raise Exception("PDF导出需要安装 reportlab 库: pip install reportlab")
        except Exception as e:
            raise Exception(f"PDF导出失败: {str(e)}")


if __name__ == "__main__":
    # 测试代码
    print("结果导出模块已加载")
